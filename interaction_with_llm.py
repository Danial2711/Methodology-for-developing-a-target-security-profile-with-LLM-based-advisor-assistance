import pdfplumber
import anthropic
import os
import json
import pypandoc
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from dotenv import load_dotenv

PROFILE_IDS_FILE = "profile_file_ids.json"
BASELINE_PROFILES_FOLDER = "Baseline_Profiles"

MAX_TOKENS_QA = 4096
MAX_TOKENS_RECOMMENDATIONS = 8192
MAX_TOKENS_PROFILE_SELECT = 100
MAX_TOKENS_PROFILE = 30000


def get_claude_client():
    load_dotenv()
    return anthropic.Anthropic(api_key=os.environ.get("CLAUDE_API_KEY"))


def read_system_specs():
    file_path = "User_input_docs/System specifications.pdf"
    if not os.path.exists(file_path):
        return "Специфікація системи не знайдена."
    text = ""
    with pdfplumber.open(file_path) as pdf:
        for page in pdf.pages:
            extracted = page.extract_text()
            if extracted:
                text += extracted + "\n"
    return text


def _load_profile_file_ids():
    if os.path.exists(PROFILE_IDS_FILE):
        with open(PROFILE_IDS_FILE, "r", encoding="utf-8") as f:
            return json.load(f)
    return {}


def _save_profile_file_ids(data):
    with open(PROFILE_IDS_FILE, "w", encoding="utf-8") as f:
        json.dump(data, f, indent=2, ensure_ascii=False)


def _verify_file_exists_on_api(client, file_id):
    try:
        client.beta.files.retrieve(file_id)
        return True
    except Exception:
        return False


def upload_baseline_profiles(force=False):
    if not os.path.exists(BASELINE_PROFILES_FOLDER):
        print(f"Папку '{BASELINE_PROFILES_FOLDER}' не знайдено.")
        return {}

    pdf_files = [f for f in os.listdir(BASELINE_PROFILES_FOLDER) if f.lower().endswith(".pdf")]
    if not pdf_files:
        print(f"У папці '{BASELINE_PROFILES_FOLDER}' немає PDF-файлів.")
        return {}

    client = get_claude_client()
    cached = {} if force else _load_profile_file_ids()

    for filename in pdf_files:
        existing = cached.get(filename)
        if existing and not force and _verify_file_exists_on_api(client, existing["file_id"]):
            continue

        path = os.path.join(BASELINE_PROFILES_FOLDER, filename)
        print(f"Завантажую '{filename}'...")
        with open(path, "rb") as f:
            uploaded = client.beta.files.upload(file=(filename, f, "application/pdf"))
        cached[filename] = {"file_id": uploaded.id, "filename": filename}

    _save_profile_file_ids(cached)
    return cached


def list_baseline_profiles():
    return _load_profile_file_ids()


def select_baseline_profile_auto(profiles, system_specs):
    if not profiles:
        return None, None

    client = get_claude_client()
    profile_list = "\n".join(f"- {name}" for name in profiles)

    prompt = (
        "Ти — експерт з інформаційної безпеки. "
        "На основі специфікації системи визнач, який базовий профіль безпеки найбільш підходить.\n\n"
        f"Доступні профілі:\n{profile_list}\n\n"
        f"Специфікація системи (перші 3000 символів):\n{system_specs[:3000]}\n\n"
        "Відповідай ТІЛЬКИ точною назвою файлу профілю з наведеного списку. Без пояснень. "
        'Приклад: "Profile_Level2.pdf"'
    )

    message = client.messages.create(
        model="claude-haiku-4-5-20251001",
        max_tokens=MAX_TOKENS_PROFILE_SELECT,
        messages=[{"role": "user", "content": prompt}],
    )
    selected_name = message.content[0].text.strip().strip('"').strip("'")

    if selected_name in profiles:
        return selected_name, profiles[selected_name]["file_id"]

    for name in profiles:
        if name.lower() in selected_name.lower() or selected_name.lower() in name.lower():
            return name, profiles[name]["file_id"]

    first_name, first_data = next(iter(profiles.items()))
    return first_name, first_data["file_id"]


def llm_query_qa(context, question):
    client = get_claude_client()

    prompt = f"""
Ти — дослідник у галузі інформаційної безпеки та експерт із технічного захисту інформації.
Твоє завдання — відповідати на загальні та теоретичні запитання щодо принципів безпеки, спираючись на наданий контекст нормативних документів (зокрема НД ТЗІ).

<instructions>
1. ОПИРАЙСЯ ТІЛЬКИ НА КОНТЕКСТ: Використовуй виключно інформацію з тегів <context>.
2. ПРАВИЛО ВІДСУТНОСТІ ДАНИХ: Якщо в <context> немає відповіді, так і скажи.
3. АНАЛІТИЧНІСТЬ: Пояснюй складні концепції зрозуміло, але з дотриманням технічної та академічної термінології.
</instructions>

<context>
{context}
</context>

Запитання користувача:
{question}
"""

    message = client.messages.create(
        model="claude-haiku-4-5-20251001",
        max_tokens=MAX_TOKENS_QA,
        messages=[{"role": "user", "content": [{"type": "text", "text": prompt}]}],
    )
    return message.content[0].text


def llm_query_recommendations(context, question):
    client = get_claude_client()
    pdf_data = read_system_specs()

    prompt = f"""
Ти — аудитор систем управління інформаційною безпекою та фахівець із ТЗІ.
Твоя мета — надавати рекомендації щодо впровадження конкретних КОНТРОЛІВ БЕЗПЕКИ для автоматизованої системи.

<instructions>
1. ВИКОРИСТАННЯ КОНТРОЛІВ: Аналізуй лише ті контролі безпеки, що наведені в <context>.
2. АДАПТИВНІСТЬ ДО СИСТЕМИ: Уважно вивчи архітектуру в <system_info> та поясни, як саме технічно реалізувати вибраний контроль (наприклад, [AC-2] або [IA-5]) саме для цієї системи.
3. ОБҐРУНТУВАННЯ: Для кожної рекомендації надавай пояснення щодо зменшення відповідного ризику чи загрози.
4. ФОРМАТУВАННЯ: Обов'язково вказуй ідентифікатор контролю в квадратних дужках.
</instructions>

<context>
{context}
</context>

<system_info>
{pdf_data}
</system_info>

Запит на рекомендацію:
{question}
"""

    message = client.messages.create(
        model="claude-haiku-4-5-20251001",
        max_tokens=MAX_TOKENS_RECOMMENDATIONS,
        messages=[{"role": "user", "content": [{"type": "text", "text": prompt}]}],
    )
    return message.content[0].text


def llm_query_profile(context, question, baseline_file_id=None):
    client = get_claude_client()
    pdf_data = read_system_specs()

    prompt = f"""
Ти — провідний архітектор безпеки та експерт з проєктування систем захисту інформації згідно з методологіями НД ТЗІ.
Твоє завдання — розробити ЦІЛЬОВИЙ ПРОФІЛЬ БЕЗПЕКИ, застосовуючи до КОЖНОГО контролю чітко визначену процедуру ухвалення рішення.

<decision_procedure>
Для кожного контролю базового профілю послідовно виконай чотири кроки.

Крок 1. Релевантність (ρ): чи застосовний контроль до цієї системи за <system_info>?
  Значення: applicable | partially applicable | not applicable.

Крок 2. Достатність (α): чи достатньо мінімальних базових параметрів контролю для цієї системи з урахуванням її ризиків?
  Значення: sufficient | insufficient | unknown.

Крок 3. Рішення (d) — обери РІВНО ОДНЕ значення, перевіряючи умови згори вниз:
  - Keep — якщо ρ=applicable і α=sufficient: контроль і параметри лишаються без змін;
  - Refine — якщо ρ=applicable і α=insufficient: уточни/посиль параметри (терміни, частоти, пороги) під цю систему;
  - Enhance — якщо ρ=applicable, контекст системи вказує на підвищений ризик і для контролю існує застосовне розширення: додай це розширення;
  - Compensate — якщо ρ=applicable, але контроль неможливо впровадити безпосередньо: заміни його компенсаційним заходом еквівалентного призначення;
  - Add — якщо контекст системи потребує заходу, відсутнього в базовому профілі: додай такий захід;
  - якщо ρ=not applicable: познач зміст клітинки як «Не застосовно» з коротким поясненням.

Крок 4. Джерело (src): для кожного рішення, відмінного від Keep, визнач нормативну підставу — ідентифікатор контролю та/або розділ із <context>, на якому ґрунтується рішення.
</decision_procedure>

<instructions>
1. СТРУКТУРА — СУВОРО ЗА БАЗОВИМ ПРОФІЛЕМ:
   Вивід — та сама таблиця, що в базовому профілі: ті самі стовпці, їх назви, порядок і кількість рядків, те саме форматування. Відтвори КОЖЕН рядок; жодного не пропускай.
   Зміни допустимі лише всередині клітинок і лише згідно з рішенням d із <decision_procedure>:
     - Keep — скопіюй оригінальний текст клітинки дослівно, без жодних змін;
     - Refine / Enhance — онови текст вимоги (не більше 2 речень);
     - Compensate — опиши компенсаційний захід (не більше 2 речень);
     - Add — додай новий рядок із відповідним ідентифікатором;
     - not applicable — «Не застосовно: <одне речення причини>».

2. ТОЧНІСТЬ ІДЕНТИФІКАТОРІВ: кожен рядок зберігає свій оригінальний ідентифікатор контролю.

3. ОПОРА НА КОНТЕКСТ: аналізуй лише контролі з <context> та архітектуру із <system_info>; не додавай вимог, відсутніх у джерелах.

4. ПРИМІТКИ ПІСЛЯ ТАБЛИЦІ:
   Після таблиці додай єдиний розділ «## Примітки до адаптації профілю».
   Включай ТІЛЬКИ контролі з рішенням, відмінним від Keep. Формат кожного рядка:
   - [ІДЕНТИФІКАТОР] (Рішення) — одне речення: що змінено і чому саме для цієї системи.
   Більше жодних розділів після приміток.

5. ФОРМАТ ВИВОДУ: Markdown. Тільки таблиця + примітки. Без вступу й підсумку.
</instructions>

<context>
{context}
</context>

<system_info>
{pdf_data}
</system_info>

Завдання:
{question}
"""

    user_content = []

    if baseline_file_id:
        user_content.append({
            "type": "document",
            "source": {"type": "file", "file_id": baseline_file_id},
            "title": "Базовий профіль безпеки",
            "context": (
                "Це базовий профіль безпеки. "
                "Цільовий профіль повинен відтворити його таблицю повністю — кожен рядок, "
                "кожен стовпець, те саме форматування."
            ),
        })

    user_content.append({"type": "text", "text": prompt})

    extra_headers = {"anthropic-beta": "files-api-2025-04-14"} if baseline_file_id else {}

    print("Профіль генерується, зачекайте...")
    message = client.messages.create(
        model="claude-opus-4-8",
        max_tokens=MAX_TOKENS_PROFILE,
        messages=[{"role": "user", "content": user_content}],
        extra_headers=extra_headers,
    )

    return message.content[0].text


def _add_table_borders(doc):
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                tc = cell._tc
                tcPr = tc.get_or_add_tcPr()
                tcBorders = OxmlElement('w:tcBorders')
                for side in ('top', 'left', 'bottom', 'right'):
                    el = OxmlElement(f'w:{side}')
                    el.set(qn('w:val'), 'single')
                    el.set(qn('w:sz'), '4')
                    el.set(qn('w:space'), '0')
                    el.set(qn('w:color'), '000000')
                    tcBorders.append(el)
                tcPr.append(tcBorders)


def save_profile_as_docx(markdown_text, base_filename):
    output_path = f"{base_filename}.docx"
    pypandoc.convert_text(markdown_text, "docx", format="md",
                          outputfile=output_path, extra_args=["--standalone"])
    doc = Document(output_path)
    _add_table_borders(doc)
    doc.save(output_path)
    return output_path